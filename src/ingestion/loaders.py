from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from src.utils.helpers import iter_supported_files

_log = logging.getLogger(__name__)

_TESSERACT_DEFAULT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def _configure_tesseract() -> None:
    import os
    import pytesseract
    cmd = os.environ.get("TESSERACT_CMD") or _TESSERACT_DEFAULT
    pytesseract.pytesseract.tesseract_cmd = cmd
    if not Path(cmd).is_file():
        _log.warning("[OCR] Tesseract executable not found at %s — set TESSERACT_CMD env var", cmd)


_configure_tesseract()


def _ocr_bytes(data: bytes) -> str:
    """OCR for clean/typed embedded images (PDF pages, DOCX blobs). No preprocessing."""
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(io.BytesIO(data))).strip()
    except Exception:
        return ""


def _ocr_handwritten(data: bytes, source_name: str = "") -> str:
    """OCR tuned for handwritten notebook photos: grayscale → resize 2x → autocontrast → threshold."""
    _log.debug("[OCR] image file detected: %s", source_name)
    try:
        import pytesseract
        from PIL import Image, ImageOps

        img = Image.open(io.BytesIO(data)).convert("L")          # grayscale
        img = img.resize((img.width * 2, img.height * 2), Image.LANCZOS)  # upscale
        img = ImageOps.autocontrast(img)                          # stretch contrast
        img = img.point(lambda p: 255 if p > 140 else 0)         # binarize

        text = pytesseract.image_to_string(img, config="--psm 6 --oem 3").strip()
        _log.debug("[OCR] %s → %d chars extracted", source_name, len(text))
        _log.debug("[OCR] preview: %r", text[:300])
        return text
    except Exception as exc:
        _log.warning("[OCR] failed for %s: %s", source_name, exc)
        return ""


@dataclass
class RawDocument:
    doc_id: str
    title: str
    source_path: str
    text: str
    metadata: Dict[str, str]


class DocumentLoader:
    def load_directory(self, docs_dir: str) -> List[RawDocument]:
        docs: List[RawDocument] = []
        for file_path in iter_supported_files(docs_dir):
            text = self.load_file(file_path)
            docs.append(
                RawDocument(
                    doc_id=file_path.stem,
                    title=file_path.stem.replace("_", " ").title(),
                    source_path=str(file_path),
                    text=text,
                    metadata={"extension": file_path.suffix.lower()},
                )
            )
        return docs

    def load_file(self, path: str | Path) -> str:
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8")
        if suffix in {".html", ".htm"}:
            return self._read_html(path)
        if suffix == ".pdf":
            return self._read_pdf(path)
        if suffix == ".docx":
            return self._read_docx(path)
        if suffix in {".jpg", ".jpeg"}:
            return _ocr_handwritten(path.read_bytes(), path.name)
        raise ValueError(f"Unsupported file type: {suffix}")

    def _read_html(self, path: Path) -> str:
        soup = BeautifulSoup(path.read_text(encoding="utf-8"), "lxml")
        for tag in soup(["script", "style"]):
            tag.decompose()
        lines = [line.strip() for line in soup.get_text("\n").splitlines()]
        return "\n".join([line for line in lines if line])

    def _read_pdf(self, path: Path) -> str:
        reader = PdfReader(str(path))
        texts = []
        for page in reader.pages:
            texts.append(page.extract_text() or "")
            for img in page.images:
                ocr = _ocr_bytes(img.data)
                if ocr:
                    texts.append(ocr)
        return "\n".join(texts)

    def _read_docx(self, path: Path) -> str:
        doc = DocxDocument(str(path))
        parts = [para.text for para in doc.paragraphs]
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                ocr = _ocr_bytes(rel.target_part.blob)
                if ocr:
                    parts.append(ocr)
        return "\n".join(parts)
